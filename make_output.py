#!/usr/bin/env python
import argparse
import pandas as pd
import shutil
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Series, Reference, PieChart
from openpyxl.chart.label import DataLabelList

def generate_output(template, dt):

    autoshuttles = ['AS01', 'AS02', 'AS03', 'AS04']

    outfile = '{}.xlsx'.format(dt)
    shutil.copy(template, outfile)

    book = load_workbook(outfile)
    writer = pd.ExcelWriter(outfile, engine="openpyxl")
    writer.book = book

    df_stats = pd.read_csv("{}_stats.csv".format(dt))
    df_hourly = pd.read_csv("{}_hourly.csv".format(dt))

    df_stats.to_excel(writer, sheet_name="Stats")
    df_hourly.to_excel(writer, sheet_name="Hourly")

    for autoshuttle in autoshuttles:
        details = pd.read_csv("{}_{}_details.csv".format(dt, autoshuttle))
        details.to_excel(writer, sheet_name=autoshuttle)

    hourly_chart = generate_hourly_chart(df_hourly, book)
    availability_chart = generate_availability_chart(book)
    picks_chart = generate_picks_chart(book)

    worksheet = book['Tableau']
    worksheet.add_chart(hourly_chart, 'A18')
    worksheet.add_chart(availability_chart, 'F18')
    worksheet.add_chart(picks_chart, 'D1')

    writer.save()
    writer.close()

    print("Generated file: {}".format(outfile))


def generate_hourly_chart(df_hourly, workbook):

    num_rows = df_hourly.shape[0]
    num_columns = df_hourly.shape[1]+1
    print("Generating hourly chart for {} rows and {} columns".format(num_rows, num_columns))

    chart = BarChart()
    chart.type = "col"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.style = 12
    chart.title = "Hourly"
    chart.y_axis.title = "Picks"
    chart.x_axis.title = "Hour"

    datasheet = workbook['Hourly']
    data = Reference(datasheet, min_col=3, max_col=num_columns, min_row=1, max_row=num_rows)
    titles = Reference(datasheet, min_col=2, min_row=2, max_row=num_rows)
    chart.add_data(data=data, titles_from_data=True)
    chart.set_categories(titles)

    return chart

def generate_availability_chart(workbook):
    print("Generating availability chart")

    chart = BarChart()
    chart.type = "bar"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.style = 12
    chart.title = "Availability Details"
    chart.y_axis.title = "Hours"

    datasheet = workbook['Summary']
    data = Reference(datasheet, min_col=2, max_col=3, min_row=13, max_row=17)
    titles = Reference(datasheet, min_col=1, min_row=14, max_row=17)
    chart.add_data(data=data, titles_from_data=True)
    chart.set_categories(titles)

    return chart

def generate_picks_chart(workbook):
    print("Generating picks chart")

    chart = PieChart()
    chart.title = "Picks"

    datasheet = workbook['Summary']
    data = Reference(datasheet, min_col=4, max_col=4, min_row=13, max_row=17)
    titles = Reference(datasheet, min_col=1, min_row=14, max_row=17)
    chart.add_data(data=data, titles_from_data=True)
    chart.set_categories(titles)

    chart.dataLabels = DataLabelList()
    chart.dataLabels.showPercent = True
    chart.dataLabels.showVal = True
    chart.dataLabels.showCatName = False
    chart.dataLabels.showSerName = False
    chart.dataLabels.showLegendKey = False

    return chart


if __name__ == '__main__':
    parser = argparse.ArgumentParser()

    parser.add_argument("-d", "--date", required=True)
    parser.add_argument("-t", "--template", default="template.xlsx")

    args = parser.parse_args()

    generate_output(args.template, args.date)
